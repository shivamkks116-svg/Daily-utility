"""Tests for POST /api/docx/to-pdf endpoint."""
import base64
import io
import os

import pytest
import requests
from docx import Document

BASE_URL = os.environ.get("EXPO_PUBLIC_BACKEND_URL", "https://daily-utility-ai.preview.emergentagent.com").rstrip("/")
ENDPOINT = f"{BASE_URL}/api/docx/to-pdf"


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE_URL}/api/auth/guest", json={"name": "TEST_docx"}, timeout=15)
    assert r.status_code == 200, f"guest auth failed: {r.status_code} {r.text}"
    body = r.json()
    tok = body.get("session_token") or body.get("token") or body.get("access_token")
    assert tok, f"No token in guest response: {body}"
    return tok


@pytest.fixture(scope="module")
def sample_docx_b64():
    """Build a small in-memory DOCX with a heading and paragraphs."""
    doc = Document()
    doc.add_heading("Test Document", level=1)
    doc.add_paragraph("This is the first paragraph of a small DOCX file used for testing conversion.")
    doc.add_paragraph("Second paragraph contains some more text so the resulting PDF is non-trivial.")
    doc.add_heading("A subheading", level=2)
    doc.add_paragraph("Final paragraph with additional text for word wrapping validation.")
    buf = io.BytesIO()
    doc.save(buf)
    return base64.b64encode(buf.getvalue()).decode("ascii")


class TestDocxToPdfHappyPath:
    def test_convert_docx_returns_pdf(self, token, sample_docx_b64):
        headers = {"Authorization": f"Bearer {token}"}
        body = {"file_base64": sample_docx_b64, "title": "test.docx"}
        r = requests.post(ENDPOINT, json=body, headers=headers, timeout=60)
        assert r.status_code == 200, f"Expected 200, got {r.status_code}: {r.text[:400]}"
        data = r.json()
        assert "file_base64" in data and isinstance(data["file_base64"], str) and data["file_base64"]
        assert "size" in data and isinstance(data["size"], int) and data["size"] > 500
        assert data.get("pages", 0) >= 1

        pdf_bytes = base64.b64decode(data["file_base64"])
        assert pdf_bytes.startswith(b"%PDF-"), f"Not a PDF, starts with: {pdf_bytes[:8]!r}"
        assert len(pdf_bytes) == data["size"]

        # Optional: open with fitz for real page count
        try:
            import fitz  # noqa
            pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
            assert pdf_doc.page_count >= 1
            pdf_doc.close()
        except ImportError:
            pass


class TestDocxToPdfErrors:
    def test_no_auth_returns_401(self, sample_docx_b64):
        r = requests.post(ENDPOINT, json={"file_base64": sample_docx_b64}, timeout=30)
        assert r.status_code == 401, f"Expected 401, got {r.status_code}: {r.text[:200]}"

    def test_bad_token_returns_401(self, sample_docx_b64):
        r = requests.post(
            ENDPOINT,
            json={"file_base64": sample_docx_b64},
            headers={"Authorization": "Bearer invalidtoken"},
            timeout=30,
        )
        assert r.status_code == 401, f"Expected 401, got {r.status_code}: {r.text[:200]}"

    def test_missing_file_field_returns_422(self, token):
        r = requests.post(
            ENDPOINT,
            json={"title": "no file"},
            headers={"Authorization": f"Bearer {token}"},
            timeout=30,
        )
        assert r.status_code == 422, f"Expected 422, got {r.status_code}: {r.text[:200]}"

    def test_invalid_base64_returns_400(self, token):
        # Contains characters not valid in base64 alphabet
        r = requests.post(
            ENDPOINT,
            json={"file_base64": "!!!not@@@base64###"},
            headers={"Authorization": f"Bearer {token}"},
            timeout=30,
        )
        # Server may accept lenient b64 decode; then the downstream DOCX open fails with 400.
        assert r.status_code == 400, f"Expected 400, got {r.status_code}: {r.text[:200]}"

    def test_non_docx_bytes_returns_400(self, token):
        b = base64.b64encode(b"hello world, this is not a docx").decode("ascii")
        r = requests.post(
            ENDPOINT,
            json={"file_base64": b},
            headers={"Authorization": f"Bearer {token}"},
            timeout=30,
        )
        assert r.status_code == 400, f"Expected 400, got {r.status_code}: {r.text[:200]}"
        assert "Could not open DOCX" in r.text or "docx" in r.text.lower()
