"""Tests for POST /api/docx/read and regression for /api/docx/to-pdf."""
import base64
import io
import os

import pytest
import requests
from docx import Document

BASE_URL = os.environ["EXPO_PUBLIC_BACKEND_URL"].rstrip("/")
READ_ENDPOINT = f"{BASE_URL}/api/docx/read"
TO_PDF_ENDPOINT = f"{BASE_URL}/api/docx/to-pdf"


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE_URL}/api/auth/guest", json={"name": "tester"}, timeout=15)
    assert r.status_code == 200, f"guest auth failed: {r.status_code} {r.text}"
    body = r.json()
    tok = body.get("session_token")
    assert tok, f"No token in guest response: {body}"
    return tok


@pytest.fixture(scope="module")
def structured_docx_b64():
    """DOCX with the exact structure required by the review request."""
    doc = Document()
    doc.add_heading("Chapter 1", level=1)
    doc.add_paragraph("Body text 1")
    doc.add_heading("Section 1.1", level=2)
    doc.add_paragraph("Body text 2")
    buf = io.BytesIO()
    doc.save(buf)
    return base64.b64encode(buf.getvalue()).decode("ascii")


@pytest.fixture(scope="module")
def multi_para_docx_b64():
    """A DOCX with several paragraphs to force a multi-line PDF."""
    doc = Document()
    doc.add_heading("Test Document", level=1)
    doc.add_paragraph("This is the first paragraph of a small DOCX file used for testing conversion.")
    doc.add_paragraph("Second paragraph with additional content so the PDF has real body text.")
    doc.add_heading("A subheading", level=2)
    doc.add_paragraph("Final paragraph with additional text for word wrapping validation.")
    buf = io.BytesIO()
    doc.save(buf)
    return base64.b64encode(buf.getvalue()).decode("ascii")


# ---------------------- /api/docx/read ----------------------
class TestDocxReadHappyPath:
    def test_read_returns_expected_blocks(self, token, structured_docx_b64):
        headers = {"Authorization": f"Bearer {token}"}
        body = {"file_base64": structured_docx_b64}
        r = requests.post(READ_ENDPOINT, json=body, headers=headers, timeout=30)
        assert r.status_code == 200, f"Expected 200, got {r.status_code}: {r.text[:400]}"
        data = r.json()

        assert "blocks" in data and isinstance(data["blocks"], list)
        assert data.get("paragraph_count") == 4, f"paragraph_count mismatch: {data.get('paragraph_count')}"
        assert data.get("char_count", 0) > 0, f"char_count should be > 0: {data.get('char_count')}"

        blocks = data["blocks"]
        # Block 0: Chapter 1 heading level 1
        assert blocks[0]["text"] == "Chapter 1", f"blocks[0].text = {blocks[0]['text']!r}"
        assert blocks[0]["level"] == 1, f"blocks[0].level = {blocks[0]['level']}"
        assert blocks[0]["bold"] is True, f"blocks[0].bold = {blocks[0]['bold']}"

        # Block 1: Body text 1
        assert blocks[1]["text"] == "Body text 1"
        assert blocks[1]["level"] == 0
        assert blocks[1]["bold"] is False

        # Block 2: Section 1.1 heading level 2
        assert blocks[2]["text"] == "Section 1.1"
        assert blocks[2]["level"] == 2
        assert blocks[2]["bold"] is True

        # Block 3: Body text 2
        assert blocks[3]["text"] == "Body text 2"
        assert blocks[3]["level"] == 0
        assert blocks[3]["bold"] is False


class TestDocxReadErrors:
    def test_no_auth_returns_401(self, structured_docx_b64):
        r = requests.post(READ_ENDPOINT, json={"file_base64": structured_docx_b64}, timeout=15)
        assert r.status_code == 401, f"Expected 401, got {r.status_code}: {r.text[:200]}"

    def test_bad_token_returns_401(self, structured_docx_b64):
        r = requests.post(
            READ_ENDPOINT,
            json={"file_base64": structured_docx_b64},
            headers={"Authorization": "Bearer invalidtoken"},
            timeout=15,
        )
        assert r.status_code == 401, f"Expected 401, got {r.status_code}: {r.text[:200]}"

    def test_missing_field_returns_422(self, token):
        r = requests.post(
            READ_ENDPOINT,
            json={},
            headers={"Authorization": f"Bearer {token}"},
            timeout=15,
        )
        assert r.status_code == 422, f"Expected 422, got {r.status_code}: {r.text[:200]}"

    def test_invalid_base64_returns_400(self, token):
        r = requests.post(
            READ_ENDPOINT,
            json={"file_base64": "!!!not@@@base64###"},
            headers={"Authorization": f"Bearer {token}"},
            timeout=15,
        )
        assert r.status_code == 400, f"Expected 400, got {r.status_code}: {r.text[:200]}"

    def test_non_docx_bytes_returns_400(self, token):
        b = base64.b64encode(b"hello world, this is not a docx").decode("ascii")
        r = requests.post(
            READ_ENDPOINT,
            json={"file_base64": b},
            headers={"Authorization": f"Bearer {token}"},
            timeout=15,
        )
        assert r.status_code == 400, f"Expected 400, got {r.status_code}: {r.text[:200]}"


# ---------------------- /api/docx/to-pdf regression ----------------------
class TestDocxToPdfRegression:
    def test_to_pdf_happy_path(self, token, multi_para_docx_b64):
        headers = {"Authorization": f"Bearer {token}"}
        body = {"file_base64": multi_para_docx_b64, "title": "test.docx"}
        r = requests.post(TO_PDF_ENDPOINT, json=body, headers=headers, timeout=60)
        assert r.status_code == 200, f"Expected 200, got {r.status_code}: {r.text[:400]}"
        data = r.json()

        assert "file_base64" in data and data["file_base64"]
        assert data.get("pages", 0) >= 1, f"pages should be >= 1, got {data.get('pages')}"

        pdf_bytes = base64.b64decode(data["file_base64"])
        assert pdf_bytes.startswith(b"%PDF-"), f"Not a PDF, starts with: {pdf_bytes[:8]!r}"

        # Verify pages field reflects the true out_doc.page_count
        try:
            import fitz
            pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
            actual_pages = pdf_doc.page_count
            pdf_doc.close()
            assert data["pages"] == actual_pages, (
                f"pages field ({data['pages']}) != actual PDF page_count ({actual_pages})"
            )
        except ImportError:
            pass
